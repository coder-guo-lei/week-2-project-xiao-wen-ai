from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

OUT = r"D:\Desktop\xiaowen\小文智能APP项目汇报.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
FONT = "Microsoft YaHei"
BG = RGBColor(8, 12, 30)
PANEL = RGBColor(24, 34, 68)
WHITE = RGBColor(245, 248, 255)
MUTED = RGBColor(176, 190, 225)
CYAN = RGBColor(58, 220, 255)
BLUE = RGBColor(74, 144, 255)
PURPLE = RGBColor(155, 92, 255)
GREEN = RGBColor(76, 230, 170)
ORANGE = RGBColor(255, 180, 90)
PINK = RGBColor(255, 92, 180)
ACC = [CYAN, BLUE, PURPLE, GREEN, ORANGE, PINK]

def add_bg(slide):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = BG; r.line.fill.background()
    for x,y,s,c,t in [(10.5,-.5,3.2,PURPLE,55),(-.8,5.3,2.7,BLUE,68),(5.2,1.1,1.4,CYAN,82)]:
        o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(s), Inches(s))
        o.fill.solid(); o.fill.fore_color.rgb = c; o.fill.transparency = t; o.line.fill.background()

def text(slide, x, y, w, h, val, size=14, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    b = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = val; p.alignment = align
    p.font.name = FONT; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
    return b

def footer(slide, n):
    text(slide, .55, 7.05, 6, .25, "小文智能 APP · AI Voice Assistant Project", 8, RGBColor(130,150,195))
    text(slide, 12.25, 7.03, .55, .25, f"{n:02d}", 10, CYAN, True, PP_ALIGN.RIGHT)

def title(slide, name, sub, n):
    add_bg(slide)
    text(slide, .7, .42, 9, .55, name, 28, WHITE, True)
    text(slide, .72, 1.02, 10.5, .36, sub, 11.5, MUTED)
    a = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.72), Inches(1.42), Inches(1.2), Inches(.05))
    a.fill.solid(); a.fill.fore_color.rgb = CYAN; a.line.fill.background()
    footer(slide, n)

def card(slide, x, y, w, h, head, body, color):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = PANEL; s.fill.transparency = 4
    s.line.color.rgb = RGBColor(72,92,145); s.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(.08), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.fill.background()
    text(slide, x+.22, y+.13, w-.4, .3, head, 13, WHITE, True)
    box = slide.shapes.add_textbox(Inches(x+.22), Inches(y+.57), Inches(w-.42), Inches(h-.67))
    tf = box.text_frame; tf.word_wrap = True
    for i, line in enumerate(body.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = FONT; p.font.size = Pt(10.2); p.font.color.rgb = MUTED

def metric(slide, x, y, v, lab, color):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.35), Inches(1.02))
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor(30,43,86); s.line.color.rgb = color; s.line.width = Pt(1.2)
    text(slide, x+.1, y+.12, 2.15, .36, v, 22, color, True, PP_ALIGN.CENTER)
    text(slide, x+.1, y+.62, 2.15, .25, lab, 9.5, MUTED, False, PP_ALIGN.CENTER)

def chip(slide, x, y, val, color):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.58), Inches(.36))
    s.fill.solid(); s.fill.fore_color.rgb = color; s.fill.transparency = 18; s.line.color.rgb = color
    text(slide, x+.05, y+.06, 1.48, .22, val, 9.5, WHITE, True, PP_ALIGN.CENTER)

def cover():
    s = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s)
    text(s,.78,1.18,7.8,.35,"AI VOICE ASSISTANT · PROJECT REPORT",12,CYAN,True)
    text(s,.75,1.68,7.8,.82,"小文智能 APP 项目汇报",38,WHITE,True)
    text(s,.8,2.58,7.5,.55,"集 AI 对话、语音交互、天气查询、音乐播放、文生图、划词翻译与桌面控制于一体的智能助手原型",16,MUTED)
    for i,(v,c) in enumerate([("React + Vite",BLUE),("Flask API",GREEN),("DashScope",PURPLE),("Web Speech",CYAN)]): chip(s,.82+i*1.82,3.38,v,c)
    text(s,.82,6.25,4,.3,"汇报人：刘汉文    日期：2026",11,RGBColor(150,166,210)); footer(s,1)

def agenda():
    s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"目录","从项目定位到落地实现，再到问题复盘与后续规划",2)
    items=["项目定位与应用场景","应用方式、方法与核心数据","系统架构与技术实现","工作亮点","问题与改进","下季度计划与感谢"]
    for i,it in enumerate(items):
        y=1.65+i*.82; text(s,.95,y,.55,.3,f"{i+1:02d}",15,CYAN,True,PP_ALIGN.CENTER); text(s,1.65,y,6,.3,it,16,WHITE,True)

def positioning():
    s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"项目定位：面向个人桌面的 AI 智能助手","通过自然语言和语音交互，把 AI 能力转化为日常可用的应用入口",3)
    card(s,.75,1.7,3.75,1.55,"用户侧价值","一句话完成查询、播放、生成、翻译和打开应用\n降低操作门槛，强化语音助手体验",CYAN)
    card(s,4.8,1.7,3.75,1.55,"产品侧价值","将多个 AI 与生活服务能力汇聚到统一界面\n形成可演示、可扩展、可迭代的 App 原型",PURPLE)
    card(s,8.85,1.7,3.75,1.55,"技术侧价值","覆盖前端组件化、后端 API、模型调用、异步任务\n适合作为课程设计、项目展示和 AI 应用样板",GREEN)
    text(s,.85,4.0,11.8,.36,"核心定位：不是单一聊天机器人，而是一个可执行任务的 AI 智能应用入口。",19,WHITE,True,PP_ALIGN.CENTER)
    for i,(h,b,c) in enumerate([("生活助手","天气 / 音乐 / 问答",BLUE),("创作助手","AI 对话 / 文生图",PURPLE),("语音助手","识别 / 唤醒 / 朗读",CYAN),("桌面助手","打开应用 / 快捷操作",GREEN)]): card(s,1.15+i*2.9,4.78,2.35,.95,h,b,c)

def data():
    s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"应用方式、实现方法与核心数据","用户一句自然语言指令，系统自动完成意图识别、能力调用和结果展示",4)
    for i,(h,b,c) in enumerate([("输入","文字 / 语音 / 示例指令",CYAN),("识别","关键词规则 + 模式判断",BLUE),("调用","AI / 天气 / 音乐 / 图片 / 本机应用",PURPLE),("反馈","卡片展示 + 日志 + 朗读",GREEN)]): card(s,.85+i*3.05,1.62,2.55,1.12,h,b,c)
    metric(s,.95,3.45,"12+","已实现核心功能点",CYAN); metric(s,3.65,3.45,"3","后端核心 API",GREEN); metric(s,6.35,3.45,"10+","前端组件模块",PURPLE); metric(s,9.05,3.45,"12","最近对话上下文条数",ORANGE)
    card(s,.95,5.08,3.75,1.1,"前端实现","React + Vite\n组件化界面 + localStorage\nWeb Speech API 语音识别与朗读",BLUE)
    card(s,4.9,5.08,3.75,1.1,"后端实现","Flask + Flask-CORS\n统一 /api/send-task 指令入口\nrequests 调用第三方服务",GREEN)
    card(s,8.85,5.08,3.55,1.1,"AI 实现","DashScope 对话模型\nDashScope 文生图任务\n多轮上下文 messages",PURPLE)

def architecture():
    s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"系统架构：前后端分离 + 多服务能力编排","以统一指令入口连接多个 AI 与工具服务",5)
    blocks=[("用户交互层","文字输入\n语音识别\n唤醒词\n示例指令\n划词工具",CYAN),("前端展示层","ChatPanel\nWeatherCard\nMusicPlayer\nImagePreview\nLogPanel / ModeBar",BLUE),("后端调度层","Flask API\nparse_command\n意图路由\n上下文管理\n白名单应用启动",GREEN),("外部能力层","DashScope 大模型\nDashScope 文生图\n高德天气\n网易云搜索\nWindows 应用",PURPLE)]
    for i,(h,b,c) in enumerate(blocks): card(s,.75+i*3.1,1.65,2.62,3.85,h,b,c)
    text(s,1.0,6.15,11.5,.38,"设计思路：前端负责体验和展示，后端负责安全调度与第三方能力调用，统一返回 type 字段驱动页面组件切换。",15,WHITE,True,PP_ALIGN.CENTER)

def highlights():
    s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"工作亮点","从“能聊天”升级为“能听、能说、能执行任务”的 AI 应用",6)
    items=[("01 多轮对话能力","支持连续追问：先问承德美食，再追问这些在哪能吃到。"),("02 AI 多能力集成","对话、文生图、翻译、天气、音乐和桌面应用统一入口。"),("03 语音助手体验","支持语音识别、唤醒词“小文小文”和回复朗读。"),("04 异步文生图体验","任务提交后进入生成中状态，通过轮询展示结果。"),("05 安全桌面控制","白名单机制打开本机应用，兼顾可用性与安全性。"),("06 组件化可扩展","前端组件和 Hook 拆分，后端统一指令路由。")]
    for i,(h,b) in enumerate(items): card(s,.8+(i%2)*6.1,1.58+(i//2)*1.55,5.65,1.18,h,b,ACC[i])

def problems():
    s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"问题与改进","当前版本已能演示核心能力，但仍需要从稳定性、智能化和工程化继续提升",7)
    rows=[("意图识别","关键词规则为主","引入大模型意图分类，减少误判"),("音乐播放","受版权和外链限制","增加多结果选择、合法音乐源和失败重试"),("语音能力","依赖浏览器支持","增加纠错、自动朗读和语速音调调节"),("后端结构","功能集中在 app.py","拆分 services、utils、config 模块"),("部署形态","本地开发模式","关闭 debug，增加生产配置和日志监控")]
    for i,(a,b,c) in enumerate(rows):
        y=1.75+i*.86; card(s,.85,y,2.3,.58,a,"",ACC[i%6]); text(s,3.45,y+.08,3.1,.25,b,11,MUTED); text(s,7.0,y+.08,5.3,.25,c,11,MUTED)

def plan():
    s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"下季度计划","围绕“更智能、更稳定、更可展示、更可扩展”推进下一阶段迭代",8)
    plans=[("第 1 阶段\n体验增强","完善多轮记忆\n自动朗读开关\n优化默认示例\n增强错误提示",CYAN),("第 2 阶段\n能力扩展","日程 / 备忘录\n图片历史与下载\n音乐收藏\n联网搜索入口",PURPLE),("第 3 阶段\n工程优化","后端服务模块化\n统一响应结构\n日志和错误码\n生产环境配置",GREEN),("第 4 阶段\n部署展示","Docker 或云部署\nHTTPS / Nginx\n演示脚本\n完善答辩材料",ORANGE)]
    for i,(h,b,c) in enumerate(plans): card(s,.8+i*3.05,1.75,2.65,3.65,h,b,c)
    text(s,1.0,6.08,11.2,.35,"目标：从“本地可运行的 AI 助手原型”升级为“可持续迭代、可部署展示的智能应用平台”。",16,WHITE,True,PP_ALIGN.CENTER)

def demo():
    s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"建议演示路径","用一条完整用户路径展示小文智能 APP 的能力闭环",9)
    demos=[("1. 语音唤醒","说“小文小文”，进入指令识别"),("2. 生活查询","询问“承德天气怎么样”"),("3. 连续对话","先问承德美食，再追问位置"),("4. 音乐播放","说“播放稻香”，展示播放器"),("5. AI 创作","说“画一只水彩小猫”"),("6. 桌面控制","说“打开计算器”")]
    for i,(h,b) in enumerate(demos): card(s,.95+(i%3)*4.05,1.65+(i//3)*2.0,3.45,1.35,h,b,ACC[i])
    text(s,1.0,6.05,11.2,.4,"演示重点：让观众看到小文不是静态页面，而是可以理解指令、调用能力并完成任务的 AI 应用。",15,WHITE,True,PP_ALIGN.CENTER)

def thanks():
    s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s)
    text(s,.9,1.65,11.5,.65,"感谢聆听",42,WHITE,True,PP_ALIGN.CENTER)
    text(s,1.45,2.55,10.4,.5,"小文智能 APP 将继续向更自然、更稳定、更具执行力的个人 AI 助手演进",18,MUTED,False,PP_ALIGN.CENTER)
    for i,(v,c) in enumerate([("能听",CYAN),("能说",BLUE),("能聊",PURPLE),("能画",PINK),("能开",GREEN)]): chip(s,2.0+i*1.85,3.55,v,c)
    text(s,3.0,5.05,7.3,.45,"Q&A",28,CYAN,True,PP_ALIGN.CENTER); footer(s,10)

for f in [cover, agenda, positioning, data, architecture, highlights, problems, plan, demo, thanks]:
    f()
prs.save(OUT)
print(OUT)
