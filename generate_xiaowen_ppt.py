"""
生成「小文智能语音助手」答辩 PPTX（与《答辩文档-小文智能语音助手.md》终稿对齐）。

依赖：pip install python-pptx

用法（仓库根目录）:
    python generate_xiaowen_ppt.py

输出：xiaowen-defense-FINAL.pptx（ASCII 文件名，避免部分终端编码问题）
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

_ROOT = Path(__file__).resolve().parent
OUT = str(_ROOT / "xiaowen-defense-FINAL.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = prs.slide_width, prs.slide_height
FONT = "Microsoft YaHei"
C = {
    "bg": RGBColor(8, 12, 30),
    "panel": RGBColor(22, 31, 62),
    "panel2": RGBColor(30, 43, 86),
    "white": RGBColor(248, 250, 255),
    "muted": RGBColor(174, 190, 225),
    "cyan": RGBColor(58, 220, 255),
    "blue": RGBColor(74, 144, 255),
    "purple": RGBColor(155, 92, 255),
    "green": RGBColor(76, 230, 170),
    "orange": RGBColor(255, 180, 90),
    "pink": RGBColor(255, 92, 180),
    "red": RGBColor(255, 105, 130),
}
ACC = [C["cyan"], C["blue"], C["purple"], C["green"], C["orange"], C["pink"]]
slide_no = [0]  # mutable counter for footer


def fill(s, color, trans=0):
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.fill.transparency = trans
    s.line.fill.background()


def line(s, color, w=1, trans=0):
    s.line.color.rgb = color
    s.line.width = Pt(w)
    s.line.transparency = trans


def txt(slide, x, y, w, h, text, size=14, color=None, bold=False, align=PP_ALIGN.LEFT):
    b = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color or C["white"]
    return b


def bg(slide):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fill(r, C["bg"])
    for x, y, sz, col, tr in [(10.4, -0.5, 3.2, C["purple"], 55), (-0.8, 5.2, 2.7, C["blue"], 65), (5.2, 1.1, 1.5, C["cyan"], 82)]:
        o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(sz), Inches(sz))
        fill(o, col, tr)
    for i in range(8):
        ln = slide.shapes.add_connector(1, Inches(0.35), Inches(0.6 + i * 0.78), Inches(13), Inches(0.6 + i * 0.78))
        line(ln, RGBColor(40, 55, 95), 0.35, 62)


def footer(slide):
    slide_no[0] += 1
    n = slide_no[0]
    txt(slide, 0.55, 7.04, 6, 0.25, "小文智能语音助手 · 答辩材料（终稿）", 8, RGBColor(130, 150, 195))
    txt(slide, 12.25, 7.02, 0.55, 0.25, f"{n:02d}", 10, C["cyan"], True, PP_ALIGN.RIGHT)


def title(slide, t, sub):
    bg(slide)
    txt(slide, 0.7, 0.42, 11.5, 0.55, t, 26, C["white"], True)
    if sub:
        txt(slide, 0.72, 1.02, 12.0, 0.38, sub, 11.5, C["muted"])
    a = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.42), Inches(1.2), Inches(0.05))
    fill(a, C["cyan"])
    footer(slide)


def card(slide, x, y, w, h, t, body, color):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(s, C["panel"], 5)
    line(s, RGBColor(70, 90, 145), 1, 35)
    b = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    fill(b, color)
    txt(slide, x + 0.22, y + 0.12, w - 0.38, 0.28, t, 13, C["white"], True)
    box = slide.shapes.add_textbox(Inches(x + 0.22), Inches(y + 0.52), Inches(w - 0.42), Inches(h - 0.62))
    tf = box.text_frame
    tf.word_wrap = True
    for i, ln in enumerate(body.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln
        p.font.name = FONT
        p.font.size = Pt(10)
        p.font.color.rgb = C["muted"]
        p.space_after = Pt(2)


def chip(slide, x, y, t, c):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.55), Inches(0.34))
    fill(s, c, 18)
    line(s, c, 1, 8)
    txt(slide, x + 0.04, y + 0.02, 1.47, 0.28, t, 9, C["white"], True, PP_ALIGN.CENTER)


# ---------- 1 封面 ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
for r, col, tr in [(2.6, C["cyan"], 70), (2.05, C["purple"], 76), (1.5, C["blue"], 82)]:
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.25 + (2.6 - r) / 2), Inches(1.75 + (2.6 - r) / 2), Inches(r), Inches(r))
    o.fill.background()
    line(o, col, 2, tr)
o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.05), Inches(2.55), Inches(1), Inches(1))
fill(o, C["cyan"], 8)
line(o, C["white"], 1, 35)
txt(s, 0.78, 1.12, 8.5, 0.32, "毕业设计 / 课程设计 · 项目答辩", 11.5, C["cyan"], True)
txt(s, 0.75, 1.58, 9.5, 0.85, "小文智能语音助手", 36, C["white"], True)
txt(
    s,
    0.8,
    2.52,
    11.5,
    0.75,
    "浏览器端 AI 助手：自然语言与语音指令 → 统一意图路由 → 对话 / 天气 / 音乐 / 文生图 / 图表 / 看图 / 翻译朗读 / 本机应用 / 模拟世界",
    14.5,
    C["muted"],
)
for i, (t, c) in enumerate(
    [
        ("React 19 + Vite 8", C["blue"]),
        ("Flask + Blueprint", C["green"]),
        ("百炼 / 高德 / 讯飞", C["purple"]),
        ("REST + JSON", C["cyan"]),
    ]
):
    chip(s, 0.78 + i * 1.78, 3.42, t, c)
txt(s, 0.82, 6.22, 5.5, 0.28, "汇报人：刘汉文    2026", 11, RGBColor(150, 166, 210))
footer(s)

# ---------- 2 目录（对齐答辩文档结构）----------
s = prs.slides.add_slide(prs.slide_layouts[6])
title(s, "目录", "与《答辩文档》章节对应，便于展开讲述")
items = [
    ("01", "总览：定位、设计思想、分层架构"),
    ("02", "技术选型：为何 React / Vite，优势何在"),
    ("03", "技术选型：为何 Python / Flask，优势何在"),
    ("04", "前后端协作与主要接口"),
    ("05", "第三方服务与工程亮点"),
    ("06", "桌面吉祥物与交互体验"),
    ("07", "产出、局限与演示建议"),
    ("08", "总结与致谢"),
]
for i, (no, t) in enumerate(items):
    y = 1.58 + i * 0.66
    txt(s, 0.88, y, 0.5, 0.28, no, 14, C["cyan"], True, PP_ALIGN.CENTER)
    txt(s, 1.55, y, 10.5, 0.3, t, 15.5, C["white"], True)
    ln = s.shapes.add_connector(1, Inches(1.48), Inches(y + 0.38), Inches(11.9), Inches(y + 0.38))
    line(ln, RGBColor(55, 73, 120), 0.5, 45)

# ---------- 3 总览：一句话 + 四大思想 ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
title(s, "项目总览", "单入口、意图驱动、前后端分离、安全白名单")
txt(s, 0.85, 1.55, 11.8, 0.95, "在浏览器中运行的「小文」：用户通过文字或语音发出自然语言指令，系统在统一意图路由 parse_command 下完成多种能力，并以 REST API + JSON 与前端联动。", 13.5, C["white"])
for i, (t, b) in enumerate(
    [
        ("单入口", "POST /api/send-task 统一接指令"),
        ("意图驱动", "按约定顺序匹配分支，映射到具体业务"),
        ("前后端分离", "密钥与第三方调用在后端 .env；前端只持 UI 状态"),
        ("安全白名单", "本机启动仅允许配置表映射，禁止用户原文进 shell"),
    ]
):
    card(s, 0.78 + (i % 2) * 6.15, 2.55 + (i // 2) * 1.38, 5.75, 1.22, t, b, ACC[i])

# ---------- 4 分层架构（文字版）----------
s = prs.slides.add_slide(prs.slide_layouts[6])
title(s, "系统架构（逻辑分层）", "表现层 → 接入层 → 业务编排 → 自建封装 / 第三方")
arch = (
    "【浏览器】React 组件 + Web Speech / getUserMedia / Audio / Selection / Geolocation\n"
    "        ↓  HTTP / JSON / multipart / 音频二进制\n"
    "【Flask】Blueprint：/api/send-task、/api/tts、/api/generate-chart 等\n"
    "        ↓  函数调用\n"
    "【task_parser】意图解析：天气、音乐、图表、对话、模拟世界、本机启动…\n"
    "        ├─ services：讯飞 TTS / IAT\n"
    "        └─ HTTP：DashScope、高德、DeepSeek（可选）"
)
box = s.shapes.add_textbox(Inches(0.82), Inches(1.58), Inches(11.7), Inches(4.85))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = arch
p.font.name = FONT
p.font.size = Pt(12.5)
p.font.color.rgb = C["muted"]
for i, (t, c) in enumerate([("前端 CSR", C["blue"]), ("REST 接入", C["cyan"]), ("意图编排", C["green"]), ("云与语音", C["purple"])]):
    chip(s, 0.85 + i * 1.72, 6.55, t, c)

# ---------- 5 技术选型：React + Vite ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
title(s, "为何选用 React 与 Vite", "答辩可复述：动机 → 在本项目中的收益")
card(
    s,
    0.75,
    1.55,
    3.95,
    2.45,
    "React",
    "多面板 contentType 切换：声明式 UI = f(状态)\n组件边界：ChatPanel / ChartPanel 等\nHook 复用：语音、音乐播放器逻辑\n生态成熟，易扩展（如 TS）",
    C["blue"],
)
card(
    s,
    4.85,
    1.55,
    3.95,
    2.45,
    "使用后的优势",
    "可维护：App.jsx 主线清晰\n与浏览器 API 生命周期契合（useEffect）\nReact 19 + Vite Fast Refresh 开发效率高",
    C["cyan"],
)
card(
    s,
    8.95,
    1.55,
    3.95,
    2.45,
    "Vite",
    "ESM 冷启动快；/api 代理到 5001\npnpm build 输出 dist，职责分离\n生产可用 VITE_API_URL 配置 apiBase.js",
    C["purple"],
)
txt(s, 0.85, 6.35, 11.6, 0.42, "诚实边界：若页面极简、无多状态，静态 HTML 亦可；本项目「多模式 + Hook」才是 React 优势场景。", 11, C["muted"])

# ---------- 6 技术选型：Python + Flask ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
title(s, "为何选用 Python 与 Flask", "快速编排多云 API；框架轻、答辩链路短")
card(
    s,
    0.75,
    1.55,
    3.95,
    2.45,
    "Python",
    "HTTP + JSON 串联百炼 / 高德 / DeepSeek\nopenpyxl、正则、csv 处理图表数据\nparse_command 分支多，可读、可讲",
    C["green"],
)
card(
    s,
    4.85,
    1.55,
    3.95,
    2.45,
    "Flask",
    "REST + jsonify + 少量二进制响应\nBlueprint 分层：routes / logic / services\nCORS 显式配置，TTS 自定义响应头可读",
    C["blue"],
)
card(
    s,
    8.95,
    1.55,
    3.95,
    2.45,
    "使用后的优势",
    "app.py 装配；routes 只做 HTTP\n标准库 subprocess/pathlib 做本机安全启动\n生产可换 gunicorn，不绑云平台",
    C["orange"],
)
txt(s, 0.85, 6.35, 11.6, 0.42, "诚实边界：大团队全家桶可选 Django/FastAPI；本项目核心是意图编排与第三方 API。", 11, C["muted"])

# ---------- 7 前端能力与契约 ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
title(s, "前端能力与接口契约", "浏览器原生能力 + 统一 apiUrl")
card(
    s,
    0.75,
    1.52,
    5.9,
    2.55,
    "浏览器 API",
    "Web Speech：唤醒 + 单次指令\ngetUserMedia：麦克风 / 摄像头\nFetch + Blob：TTS、图表上传\nGeolocation：当地天气 / 附近推荐\nSelection：划词翻译工具条",
    C["cyan"],
)
card(
    s,
    6.85,
    1.52,
    5.65,
    2.55,
    "主要接口",
    "POST /api/send-task（task、history、location）\nPOST /api/tts（音频流）\nGET /api/image-status/:id\nPOST /api/generate-chart、analyze-image\nPOST /api/translate-selection",
    C["purple"],
)
txt(s, 0.82, 6.32, 11.7, 0.45, "开发态：Vite 将 /api 代理到本机 5001；生产通过 apiBase.js 配置 API 根地址。", 12, C["white"])

# ---------- 8 第三方服务 ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
title(s, "第三方服务", "密钥均在 backend/.env，不进入前端仓库")
for i, (t, b, c) in enumerate(
    [
        ("阿里云百炼", "对话、文生图、VL 看图、翻译等\nDASHSCOPE_API_KEY", C["purple"]),
        ("高德开放平台", "地理编码、天气预报\nAMAP_KEY", C["blue"]),
        ("讯飞开放平台", "TTS 朗读、可选 WAV 听写\n与控制台产品开通一致", C["cyan"]),
        ("DeepSeek（可选）", "OpenAI 兼容对话/翻译/知识库\n不能替代生图与看图", C["green"]),
    ]
):
    card(s, 0.78 + (i % 2) * 6.15, 1.55 + (i // 2) * 1.92, 5.85, 1.78, t, b, ACC[i])

# ---------- 9 功能模块映射（简表）----------
s = prs.slides.add_slide(prs.slide_layouts[6])
title(s, "功能与前端载体", "type 字段驱动 contentType 切换")
rows = [
    ("多轮对话", "ChatPanel", "DashScope / DeepSeek"),
    ("天气", "WeatherCard", "高德"),
    ("音乐", "MusicPlayer", "搜索外链"),
    ("文生图", "ImagePreview + 轮询", "百炼异步任务"),
    ("图表", "ChartPanel", "正则 + CSV/Excel"),
    ("看图 / 肤质", "ImageAnalyzer / FaceWellness", "百炼 VL"),
    ("本机应用", "文案 type:app", "白名单 launch"),
    ("朗读", "ChatPanel + xfyunTts", "讯飞 TTS"),
]
for i, (a, b, c) in enumerate(rows):
    y = 1.52 + i * 0.62
    txt(s, 0.85, y, 2.0, 0.28, a, 11.5, C["white"], True)
    txt(s, 2.95, y, 3.4, 0.28, b, 11, C["muted"])
    txt(s, 6.55, y, 5.9, 0.28, c, 11, C["cyan"])

# ---------- 10 工程亮点 ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
title(s, "工程亮点", "可演示、可解释、有安全边界")
highlights = [
    ("统一意图路由", "parse_command 顺序敏感，便于答辩讲分支优先级"),
    ("密钥后端化", "前端仅 apiUrl；.env 被 .gitignore 忽略"),
    ("异步文生图", "taskId 轮询 + 前端进度与超时对齐"),
    ("桌面吉祥物 XiaowenBot", "内联 SVG + Pointer 拖拽 + rAF 惯性 + CSS 表情\n纯前端、不调接口"),
    ("无障碍与体验", "prefers-reduced-motion 关闭甩动惯性"),
    ("质量工具", "ESLint + react-hooks；pnpm build 产物可部署"),
]
for i, (t, b) in enumerate(highlights):
    card(s, 0.78 + (i % 2) * 6.1, 1.52 + (i // 2) * 1.42, 5.85, 1.28, t, b, ACC[i % 6])

# ---------- 11 产出与局限 ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
title(s, "产出物与局限（诚实说明）", "答辩加分：主动说明边界")
card(
    s,
    0.75,
    1.52,
    5.85,
    2.55,
    "产出物",
    "pnpm dev + python app.py 可演示\npnpm build → dist 静态资源\n.env.example 配置说明\n答辩文档 + 代码导读",
    C["green"],
)
card(
    s,
    6.75,
    1.52,
    5.95,
    2.55,
    "局限",
    "音乐外链受版权与平台策略影响\n语音识别依赖浏览器与环境权限\n模拟世界等状态在进程内存，重启丢失\ntask_parser 体量大，靠章节注释维护顺序",
    C["orange"],
)

# ---------- 12 演示建议 ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
title(s, "建议演示路径", "约 8～12 分钟可裁剪")
demos = [
    ("1", "文字天气 → WeatherCard + 模式栏"),
    ("2", "AI 对话 + 多轮一句 → history"),
    ("3", "服务端 TTS 朗读 → 非浏览器自带"),
    ("4", "文生图或图表二选一 → 工作流面板"),
    ("5", "打开计算器 → 白名单安全设计"),
    ("6", "（可选）划词翻译"),
]
for i, (no, t) in enumerate(demos):
    card(s, 0.82 + (i % 3) * 4.05, 1.55 + (i // 3) * 1.42, 3.75, 1.22, f"步骤 {no}", t, ACC[i % 6])
txt(s, 0.88, 6.15, 11.5, 0.4, "演示前检查：backend/.env 已配置百炼、高德、讯飞至少保证对话 + 天气 + 朗读其一畅通。", 12, C["muted"])

# ---------- 13 总结 ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
title(s, "总结陈述", "技术选型与架构一句话收束")
txt(
    s,
    0.82,
    1.55,
    11.6,
    4.85,
    "本项目基于 React + Vite 与 Python Flask：前端因「多面板状态与 Hook 复用」选择 React，因「热更新与 /api 代理」选择 Vite；后端因「快速编排多云 REST」选择 Python，因「轻量、路由清晰、易答辩讲述」选择 Flask。\n\n"
    "通过 RESTful API 整合百炼、高德、讯飞及可选 DeepSeek；以 parse_command 意图路由为核心，以前端组件化状态与浏览器语音/媒体 API 提供交互。\n\n"
    "体现前后端分离、密钥隔离、白名单安全与异步任务处理等工程化基本素养。",
    13.5,
    C["white"],
)

# ---------- 14 致谢 ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
txt(s, 0.85, 1.72, 11.6, 0.7, "感谢聆听", 40, C["white"], True, PP_ALIGN.CENTER)
txt(s, 1.2, 2.62, 10.9, 0.45, "小文智能语音助手 — 答辩 PPT 终稿（与文档同步）", 16, C["muted"], False, PP_ALIGN.CENTER)
for i, (t, c) in enumerate([("听", C["cyan"]), ("说", C["blue"]), ("聊", C["purple"]), ("画", C["pink"]), ("看", C["orange"]), ("控", C["green"])]):
    chip(s, 1.35 + i * 1.55, 3.45, t, c)
txt(s, 3.15, 5.1, 7.0, 0.45, "Q & A", 28, C["cyan"], True, PP_ALIGN.CENTER)
footer(s)

prs.save(OUT)
print("已生成:", OUT)
